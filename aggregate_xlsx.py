#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
業績個人票（様式2）集計スクリプト
==================================

各教員が入力した「様式2 (個人票)」Excel を集めたフォルダを読み込み、
右側の「優秀研究賞一次審査用集計表」の点数（部門別・合計・項目別）を抽出して
ランキングBIツール用の data_award.json を出力します。

前提（ヒアリングに基づく）
  - 提出形態は混在：①1ファイルに複数タブ（研究者ごと）／②1人1ファイル。両方を一括処理。
  - ファイル内に所属情報は無く氏名のみ。よって所属はフォルダ階層から取得する。
      推奨フォルダ構成:  ルート/<学部>/<学科>/*.xlsx
  - 氏名は フォーム内セル(B2) → タブ名 → ファイル名 の順で補完。
  - 点数は Excel が計算済みの値（U列）をそのまま採用。
    → 各ファイルは Excel で開いて保存された（数式が計算済みの）状態である必要があります。

使い方
  pip install openpyxl
  python aggregate_xlsx.py --dir ./提出ルート --out data_award.json
  # 1ファイルの構造を確認したいとき
  python aggregate_xlsx.py --debug "サンプル.xlsx"

※ 実データ（外部非公開）はお手元の環境でのみ処理してください。
※ 実ファイルのレイアウト差異があれば、--debug の出力を見て調整します。
"""

import argparse
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string as colidx

# 右側集計表の列
COL_ITEM = colidx("R")   # 項目
COL_SUB = colidx("S")    # 細目
COL_PT = colidx("U")     # 点数

DIVISIONS = ["論文部門", "学術部門", "アクティビティー部門"]
TOTAL_LABEL = "合計"

# 学科 → 学部 対応（東京都市大学）。未知の学科は学部空欄。
DEPT2FAC = {
    "機械工学科": "理工学部", "機械システム工学科": "理工学部", "電気電子通信工学科": "理工学部",
    "医用工学科": "理工学部", "応用化学科": "理工学部", "原子力安全工学科": "理工学部", "自然科学科": "理工学部",
    "建築学科": "建築都市デザイン学部", "都市工学科": "建築都市デザイン学部",
    "情報科学科": "情報工学部", "知能情報工学科": "情報工学部",
    "環境創生学科": "環境学部", "環境経営システム学科": "環境学部",
    "社会メディア学科": "メディア情報学部", "情報システム学科": "メディア情報学部",
    "デザイン・データ科学科": "デザイン・データ科学部",
    "都市生活学科": "都市生活学部",
    "人間科学科": "人間科学部",
}


def clean(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def to_num(v):
    """点数セルを数値化。空・エラー(#DIV/0! 等)は 0。"""
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return 0.0
    s = str(v).strip()
    try:
        return float(s)
    except ValueError:
        return 0.0  # '#DIV/0!' '#VALUE!' 等


def merged_value(ws, row, col):
    """結合セルを考慮してセル値を返す（結合内なら左上の値）。"""
    cell = ws.cell(row=row, column=col)
    if cell.value not in (None, ""):
        return cell.value
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col).value
    return cell.value


def find_label(ws, text, max_row=12, max_col=12):
    """指定テキストのセル座標(row,col)を返す。見つからなければ None。"""
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if clean(merged_value(ws, r, c)) == text:
                return (r, c)
    return None


def is_form_sheet(ws) -> bool:
    """様式2（個人票）シートかどうか判定。"""
    return find_label(ws, "教員氏名") is not None or find_label(ws, "合計", max_row=60, max_col=25) is not None


def parse_sheet(ws):
    """1シート（=1研究者）を解析して dict を返す。"""
    # --- 氏名・職位・年齢（ラベルの直下セル） ---
    def value_below(label):
        pos = find_label(ws, label)
        if not pos:
            return ""
        return clean(merged_value(ws, pos[0] + 1, pos[1]))

    name = value_below("教員氏名")
    position = value_below("職位")
    age = value_below("年齢")

    # --- 集計表（点数） ---
    divisions = {}
    total = 0.0
    items = {}
    last_row = ws.max_row
    for r in range(1, last_row + 1):
        rlab = clean(merged_value(ws, r, COL_ITEM))
        slab = clean(merged_value(ws, r, COL_SUB))
        pt_cell = ws.cell(row=r, column=COL_PT).value
        if rlab in DIVISIONS:
            divisions[rlab] = to_num(pt_cell)
            continue
        if rlab == TOTAL_LABEL:
            total = to_num(pt_cell)
            continue
        # 項目行：点数セルが数値、かつラベルがある行
        if (rlab or slab) and isinstance(pt_cell, (int, float)):
            label = " ".join(x for x in [rlab, slab] if x and x not in ("項目", "細目"))
            if label and label not in ("点数",):
                items[label] = to_num(pt_cell)

    return {
        "name": name,
        "position": position,
        "age": age,
        "total": total,
        "divisions": divisions,
        "items": items,
    }


def debug_file(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"ファイル: {path}")
    print(f"シート数: {len(wb.worksheets)}")
    for ws in wb.worksheets:
        form = is_form_sheet(ws)
        print(f"\n--- シート '{ws.title}'  (個人票判定={form}) ---")
        if not form:
            continue
        rec = parse_sheet(ws)
        print(f"  氏名: {rec['name']!r}  職位: {rec['position']!r}  年齢: {rec['age']!r}")
        print(f"  合計: {rec['total']}  部門: {rec['divisions']}")
        print(f"  項目数: {len(rec['items'])}")
        for k, v in list(rec["items"].items())[:25]:
            print(f"    [{k}] = {v}")


def affiliation_from_path(root, path):
    """ルートからの相対パスで 学部/学科 を推定。

    想定フォルダ構成（学部/学科/ファイル）:
        root/理工学部/機械工学科/○○.xlsx
    1階層しかない場合は学科とみなし DEPT2FAC で学部を補完。
    """
    rel = os.path.relpath(os.path.dirname(path), root)
    parts = [] if rel in (".", "") else rel.split(os.sep)
    if len(parts) >= 2:
        return parts[0], parts[1]          # 学部, 学科
    if len(parts) == 1:
        dept = parts[0]
        return DEPT2FAC.get(dept, ""), dept  # 学科のみ → 学部を補完
    return "", ""                           # 直下＝所属不明


def main():
    ap = argparse.ArgumentParser(description="様式2 個人票 集計")
    ap.add_argument("--dir", help="提出ルートフォルダ（配下を再帰探索）")
    ap.add_argument("--out", default="data_award.json", help="出力JSON")
    ap.add_argument("--debug", metavar="XLSX", help="1ファイルの構造を診断表示して終了")
    args = ap.parse_args()

    if args.debug:
        debug_file(args.debug)
        return
    if not args.dir:
        ap.error("--dir か --debug を指定してください")

    root = os.path.abspath(args.dir)
    # 学部/学科 フォルダ配下の xlsx を再帰探索（①タブ分割・②ファイル分割を一括処理）
    paths = []
    for dirpath, _dirs, files in os.walk(root):
        for fn in sorted(files):
            if fn.lower().endswith((".xlsx", ".xlsm")) and not fn.startswith("~$"):
                paths.append(os.path.join(dirpath, fn))
    paths.sort()
    print(f"対象ファイル {len(paths)} 件", file=sys.stderr)

    researchers = []
    warnings = []
    seen = {}  # (faculty,dept,name) -> 出現回数（重複検知）
    for path in paths:
        fac, dept = affiliation_from_path(root, path)
        fn = os.path.basename(path)
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as e:  # noqa: BLE001
            print(f"  ! 読み込み失敗 {fn}: {e}", file=sys.stderr)
            warnings.append(f"読み込み失敗: {os.path.relpath(path, root)} ({e})")
            continue
        cnt = 0
        for ws in wb.worksheets:
            if not is_form_sheet(ws):
                continue
            rec = parse_sheet(ws)
            if not rec["name"] and not rec["total"]:
                continue  # 空シートはスキップ
            # 氏名: フォーム内セル → タブ名 → ファイル名 の順で補完
            name_src = "cell"
            if not rec["name"]:
                rec["name"] = ws.title
                name_src = "tab"
            if not rec["name"] or rec["name"] == "様式2 (個人票)":
                rec["name"] = os.path.splitext(fn)[0]
                name_src = "file"
            rec["faculty"] = fac
            rec["department"] = dept
            rec["_source"] = f"{os.path.relpath(path, root)} [{ws.title}]"
            rec["_name_src"] = name_src
            if not fac and not dept:
                warnings.append(f"所属不明（フォルダ直下）: {rec['name']} ← {rec['_source']}")
            key = (fac, dept, rec["name"])
            seen[key] = seen.get(key, 0) + 1
            researchers.append(rec)
            cnt += 1
        print(f"  {os.path.relpath(path, root)} → {cnt} 名", file=sys.stderr)

    for key, n in seen.items():
        if n > 1:
            warnings.append(f"氏名重複（同姓同名/二重提出の可能性 {n}件）: {key[2]}（{key[1]}）")

    out = {
        "meta": {
            "source": "様式2 個人票（優秀研究賞一次審査用集計表）",
            "divisions": DIVISIONS,
            "count": len(researchers),
            "warnings": warnings,
        },
        "researchers": researchers,
    }
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(f"書き出し完了: {args.out}（{len(researchers)} 名）", file=sys.stderr)
    if warnings:
        print(f"\n⚠ 要確認 {len(warnings)} 件:", file=sys.stderr)
        for w in warnings:
            print(f"  - {w}", file=sys.stderr)


if __name__ == "__main__":
    main()
